#!/usr/bin/env python3
"""《KYC工作指引（2026）》附件4结构 KYC Word 生成器（结论置前 · 茅台租赁默认版）。

样例：~/Downloads/KYC/上海城投控股股份有限公司_KYC报告.docx
勿使用多章「三源核验版 / Codex 样式」作为默认输出。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "1F4E79"
GRAY = "666666"
BLACK = "000000"


def _set_run(run, size=11, bold=False, color=BLACK, font="宋体"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _margins(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, v in (("top", 60), ("start", 80), ("bottom", 60), ("end", 80)):
        n = OxmlElement(f"w:{m}")
        n.set(qn("w:w"), str(v))
        n.set(qn("w:type"), "dxa")
        tcMar.append(n)
    tcPr.append(tcMar)


class Att4Builder:
    def __init__(
        self,
        company: str,
        profile: dict[str, Any] | None,
        extras: dict[str, Any] | None,
        checked_at: str,
        sources_line: str | None = None,
    ):
        self.company = company
        self.p = profile or {}
        self.x = extras or {}
        self.checked_at = checked_at
        self.name = self.p.get("name") or company
        self.sources_line = sources_line or (
            "启信慧眼 MCP、财汇企业预警通 MCP、同花顺 iFinD、万得 Wind"
        )
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        style = self.doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(4)

    def _v(self, *keys, default="—"):
        for k in keys:
            v = self.p.get(k)
            if v not in (None, ""):
                return v
        return default

    def add_p(
        self,
        text="",
        bold=False,
        size=11,
        align=None,
        space_after=4,
        font="宋体",
        color=BLACK,
    ):
        p = self.doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        _set_run(r, size=size, bold=bold, color=color, font=font)
        return p

    def add_h(self, text: str):
        return self.add_p(text, bold=True, size=12, space_after=8, font="黑体", color=NAVY)

    def add_table(self, headers, rows):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = str(h)
            _shade(c, "D6E3F0")
            _margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                for r in p.runs:
                    _set_run(r, size=9, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                _margins(cells[i])
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        _set_run(r, size=9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return t

    def _risk_bits(self) -> tuple[str, str, str]:
        """返回 (风险级别中文, 研判结论, 准入意见)."""
        override = self.x.get("risk_level")
        nature = str(self._v("enterprise_nature", default=""))
        listed = bool(self.x.get("is_listed") or self.p.get("is_listed"))
        if override:
            level = override
        elif listed or "国有" in nature or "上市" in nature:
            level = "低"
        else:
            level = "一般关注"
        judgment = self.x.get("judgment") or (
            "一般关注" if level == "低" else "审慎关注"
        )
        admission = self.x.get("admission") or (
            "附条件同意（低风险身份认定成立；信用层面附经营与负债监测条件，最终以具体业务审批为准）。"
            if level == "低"
            else "审慎准入：须补齐评级/报告与非公开核验后再进入尽调与审批。"
        )
        return level, judgment, admission

    def build(self, out_path: Path) -> Path:
        level, judgment, admission = self._risk_bits()
        date_cn = self.checked_at[:10] if self.checked_at else datetime.now().strftime("%Y-%m-%d")
        uscc = self._v("uscc", "credit_no", "unified_social_credit_code", default="待核")
        legal = self._v("legal_representative", default="待核")
        capital = self._v("registered_capital", default="待核")
        est = self._v("establishment_date", default="待核")
        addr = self._v("registered_address", default="待核")
        nature = self._v("enterprise_nature", default="待核")
        ctrl = self._v("controller", default="待核")
        rating = self._v("latest_subject_rating", default="")
        rating_org = self._v("latest_subject_rating_organization", default="")
        rating_dt = self._v("latest_subject_rating_date", default="")
        rating_txt = (
            f"{rating}（{rating_org} {rating_dt}）".strip()
            if rating and rating != "—"
            else "公开主体评级暂缺，建议跟踪评级/债券披露或征信续核"
        )
        stock = self.x.get("stock_code") or self.p.get("stock_code") or ""
        industry = self.x.get("industry") or self.p.get("industry") or "待项目化认定"
        scope = self.x.get("scope") or self.p.get("scope") or "见工商公示经营范围"
        control_path = self.x.get("control_path") or f"公开控制关系：{ctrl}"
        chair = self.x.get("chairman") or legal
        gm = self.x.get("general_manager") or "待核"
        employees = self.x.get("employees_note") or ""
        listed_note = self.x.get("listed_note") or (
            f"是（{stock}）" if stock else ("是" if self.x.get("is_listed") else "否/待核")
        )
        bond_note = self.x.get("bond_note") or "待核（注意：历史发行只数 ≠ 余额>0 存量只数）"
        guarantee_note = self.x.get("guarantee_note") or "待专项查询/客户提供担保明细。"
        change_note = self.x.get("change_note") or "未见结构化工商变更摘要；建议归档日复核公示系统。"
        penalty_note = self.x.get("penalty_note") or "公开检索未返回可结构化主体处罚；归档日保留信用中国等留痕。"
        qixin_risk_note = self.x.get("qixin_risk_note") or "启信慧眼：失信/被执行/经营异常/严重违法见正文或待查。"
        leasing_note = self.x.get("leasing_note") or "中登网近窗未检索到或未专项查询；须正式全量补齐。"
        ar_note = self.x.get("ar_note") or leasing_note
        news_note = self.x.get("news_note") or "舆情细核可结合公告/新闻续补。"
        influence_note = self.x.get("influence_note") or (
            f"法定代表人/董事长{chair}；总经理{gm}。章程层其他有影响力自然人：待审章程后确认。"
        )
        basis = self.x.get("risk_basis") or (
            f"客户企业性质为{nature}，实际控制人/控制关系为「{ctrl}」。"
            + (f"证券代码{stock}。" if stock else "")
            + "是否符合《指引》附件1低风险/高风险直接认定，以公开核验与项目材料为准。"
        )
        negative_items = self.x.get("negative_items") or [
            f"①司法/处罚：{penalty_note}",
            f"②启信风险扫描：{qixin_risk_note}",
            f"③主体评级：{rating_txt}",
            "④人行征信与中登全量登记：本次公开KYC未覆盖，须正式渠道补充。",
        ]
        fact_summary = self.x.get("fact_summary") or (
            f"{self.name}公开身份与控制权已核验；评级 {rating_txt}。"
            "经营、债务与交易结构须结合项目方案进一步判断。"
        )
        relevance = self.x.get("relevance") or (
            "公开负面事项以客户自身经营/司法为主时，应评估与本次租赁/保理还款来源的关联性；"
            "未见失信执行不等于无信用风险。"
        )
        business_impact = self.x.get("business_impact") or (
            "身份识别维度可按公开风险等级管理；信用决策须结合征信、中登、还款来源与增信，"
            "不宜仅因国有/上市身份简化实质判断。"
        )
        controls = self.x.get("controls") or (
            "①明确还款来源与增信；②持续监测财报、诉讼、担保与评级；"
            "③补齐人行征信与中登全量登记；④大额或复杂交易提高监测强度。"
        )

        # Cover
        self.add_p(
            self.name,
            bold=True,
            size=18,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=6,
            font="黑体",
        )
        self.add_p(
            "KYC报告",
            bold=True,
            size=16,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=10,
            font="黑体",
            color=NAVY,
        )
        self.add_p(
            f"报告日期：{date_cn}　｜　数据来源：{self.sources_line}",
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=GRAY,
        )
        self.add_p(
            "说明：本报告结构参照公司《KYC工作指引（2026）》附件4。"
            "基本信息、企业融资、经营预警、监管处罚、司法诉讼、公司舆情等模块依托启信慧眼与财汇等第三方数据交叉核验；"
            "人行征信信息需另行补充纸质/电子征信报告。KYC综合结论置于报告最前。",
            size=10,
            color=GRAY,
            space_after=12,
        )

        # 1 结论
        self.add_h("1、KYC综合结论")
        self.add_p(f"（1）客户风险级别为：{level}", bold=True)
        self.add_p(f"认定依据：{basis}")
        self.add_p("（2）负面事项汇总：", bold=True)
        for item in negative_items:
            self.add_p(item)
        self.add_p("（3）风险影响研判：", bold=True)
        self.add_p(f"事实摘要：{fact_summary}")
        self.add_p(f"关联性判断：{relevance}")
        self.add_p(f"对本次业务的影响：{business_impact}")
        self.add_p(f"结论：{judgment}。", bold=True)
        self.add_p("（4）对本次业务的影响及管控建议：", bold=True)
        self.add_p(controls)
        self.add_p(f"（5）是否建议准入：{admission}", bold=True)

        # 2 证照
        self.add_h("2、客户相关证照信息与行业认定")
        self.add_p(f"企业全称：{self.name}")
        self.add_p(f"统一社会信用代码：{uscc}")
        self.add_p(f"法定代表人：{legal}　｜　董事长：{chair}　｜　总经理：{gm}")
        self.add_p(f"注册资本：{capital}　｜　成立日期：{est}　｜　登记状态：{self.x.get('status') or '存续/待核'}")
        self.add_p(f"注册地址：{addr}")
        self.add_p(
            f"企业性质：{nature}　｜　是否上市：{listed_note}　｜　是否发债：{self.x.get('has_bond') or '待核'}"
        )
        self.add_p(f"客户行业认定：{industry}")
        self.add_p(f"经营范围（摘要）：{scope}")

        # 3 征信
        self.add_h("3、人行征信信息")
        self.add_p(
            "（本次启信/财汇等公开数据未覆盖人行征信明细，须由客户或银行提供纸质/电子征信报告后补充列表。"
            "五级分类非正常类须在综合结论单独说明。）"
        )
        credit_rows = self.x.get("credit_rows") or [
            ["短期借款", "待补充", "", "", ""],
            ["长期借款", "待补充", "", "", ""],
        ]
        self.add_table(["融资类型", "金额", "期限", "五级分类", "金融机构"], credit_rows)

        # 4 基本信息
        self.add_h("4、基本信息")
        self.add_p(f"（1）工商信息：见第2节。{employees}")
        self.add_p(f"（2）实际控制人：{ctrl}")
        self.add_p(f"控制路径：{control_path}")
        self.add_p("（3）股东信息（主要股东）：")
        sh_rows = self.x.get("shareholder_rows") or [
            ["待补股东表", "—", "—", "建议以启信/财汇/iFinD/Wind交叉"]
        ]
        self.add_table(["股东名称", "持股比例", "持股数量/认缴", "股东类型/性质"], sh_rows)
        self.add_p("（4）工商变更（近3年关注）：")
        self.add_p(change_note)

        # 5 融资
        self.add_h("5、企业融资")
        self.add_p("（1）有息债务相关财务指标（年报，单位：万元）：")
        fin_headers = self.x.get("finance_headers") or [
            "报告期",
            "总资产",
            "营业总收入",
            "净利润",
            "资产负债率%",
            "货币资金",
            "短期借款",
            "长期借款",
            "应付债券",
        ]
        fin_rows = self.x.get("finance_rows") or [["待补", "—", "—", "—", "—", "—", "—", "—", "—"]]
        self.add_table(fin_headers, fin_rows)
        self.add_p(f"（2）租赁融资（中登网近窗）：{leasing_note}")
        self.add_p(f"（3）应收账款融资/质押（中登网近窗）：{ar_note}")
        self.add_p(f"（4）发债与评级及存量债券：{bond_note} 主体评级：{rating_txt}。")
        self.add_p(f"（5）对外担保：{guarantee_note}")

        # 6-8
        self.add_h("6、经营预警")
        for line in self.x.get("warning_lines") or [
            "（1）股权出质：待核/未检索到。",
            "（2）动产抵押：待核/未检索到。",
            "（3）股权冻结：待核/未检索到。",
            "（4）股票质押：待核/未检索到。",
        ]:
            self.add_p(line)

        self.add_h("7、监管处罚")
        self.add_p(penalty_note)
        self.add_p(qixin_risk_note)

        self.add_h("8、司法诉讼")
        self.add_p("（1）司法案件（检索命中）：")
        jud_rows = self.x.get("judicial_rows") or [
            ["—", "公开检索未见可结构化明细或未命中", "—", "—", "—"]
        ]
        self.add_table(
            ["最新进程日期", "案件名称/案由", "案件身份", "案号", "目前进展"],
            jud_rows,
        )
        self.add_p(self.x.get("judicial_note") or "注：诉讼信息包括未决、已决及执行中；重大案件已列示。")
        self.add_p(f"（2）被执行人/失信被执行人：{self.x.get('executed_note') or '见启信/财汇检索'}")
        self.add_p(f"（3）限制高消费：{self.x.get('limit_consume_note') or '见启信检索'}")
        self.add_p(f"（4）破产程序：{self.x.get('bankruptcy_note') or '见启信/财汇检索'}")

        # 9 舆情
        self.add_h("9、公司负面舆情")
        news_rows = self.x.get("news_rows")
        if news_rows:
            self.add_table(["新闻日期", "新闻标题", "预警分类/要点", "新闻来源"], news_rows)
        self.add_p(news_note)

        # 10-13
        self.add_h("10、国家企业信用信息公示系统、信用中国等")
        self.add_p(
            self.x.get("official_note")
            or "建议业务归档时保留查询日截屏或标准化底稿溯源说明；公开第三方库“未检索到”不等于法律上绝对不存在。"
        )
        self.add_h("11、发债企业重大信息公告（若有）")
        self.add_p(self.x.get("bond_announce_note") or "建议至少覆盖近1年定期报告与临时公告。")
        self.add_h("12、上市公司重大信息公告（若有）")
        self.add_p(self.x.get("listed_announce_note") or "非上市主体可注明不适用；上市主体持续关注定期/临时公告。")
        self.add_h("13、对客户有影响力的自然人（若有）")
        self.add_p("是否进行审查：是（主要管理人员）。")
        self.add_p(influence_note)

        self.add_p("——", align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)
        self.add_p(
            "免责声明：本报告基于启信慧眼、财汇MCP、iFinD、Wind等于查询时点返回的公开及合法取得数据编制，"
            "仅供内部KYC参考；数据可能存在滞后、窗口限制、快照口径差异或不完整。"
            "人行征信、完整中登历史登记等须以正式渠道补充。公开KYC不能替代征信、中登、合同与决议核验。",
            size=9,
            color=GRAY,
        )

        self.doc.core_properties.title = f"{self.name}KYC报告"
        self.doc.core_properties.subject = "依据KYC工作指引（2026）附件4 · auto-kyc"
        self.doc.core_properties.author = "auto-kyc"
        self.doc.core_properties.keywords = "KYC, 附件4, 启信慧眼, 财汇, iFinD, Wind, 茅台租赁"
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out_path))
        return out_path


def build_att4_report(
    company: str,
    profile: dict | None,
    extras: dict | None,
    out_path: Path,
    checked_at: str | None = None,
    sources_line: str | None = None,
) -> Path:
    checked = checked_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return Att4Builder(company, profile, extras, checked, sources_line).build(out_path)

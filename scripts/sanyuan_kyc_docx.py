#!/usr/bin/env python3
"""按《KYC工作指引（2026）》与上海城投三源核验版同结构，生成 Skill KYC Word。

样本：~/Downloads/KYC/上海城投控股股份有限公司_KYC报告_2026三源核验版.docx
生成器参考：~/Downloads/KYC/build_kyc_report.py
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
PALE = "F4F7FA"
GOLD = "C59B3C"
AMBER = "C77700"
GREEN = "2E6B3D"
GRAY = "666666"
WHITE = "FFFFFF"
BLACK = "1F1F1F"


def _shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def _margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def _borders(table, color="B7C3D0", size=6) -> None:
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)


def _cell_font(cell, size=8.4, bold=False, color=BLACK, align=None) -> None:
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = "Hiragino Sans GB"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)


def _run_font(run, size=10.2, bold=False, color=BLACK) -> None:
    run.font.name = "Hiragino Sans GB"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


class SanyuanBuilder:
    def __init__(self, company: str, profile: dict | None, extras: dict | None, checked_at: str):
        self.company = company
        self.p = profile or {}
        self.extras = extras or {}
        self.checked_at = checked_at
        self.name = self.p.get("name") or company
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(0.78)
        sec.bottom_margin = Inches(0.72)
        sec.left_margin = Inches(0.82)
        sec.right_margin = Inches(0.82)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Hiragino Sans GB"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        normal.font.size = Pt(10.2)
        for style_name, size, color, before, after in [
            ("Title", 26, NAVY, 0, 10),
            ("Heading 1", 16, NAVY, 14, 7),
            ("Heading 2", 12.5, BLUE, 10, 5),
        ]:
            st = self.doc.styles[style_name]
            st.font.name = "Hiragino Sans GB"
            st._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
            st.font.size = Pt(size)
            st.font.color.rgb = RGBColor.from_string(color)
            st.font.bold = True
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
        header = sec.header.paragraphs[0]
        header.text = "KYC 尽职调查报告  |  内部资料 · auto-kyc Skill 三源核验版"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in header.runs:
            _run_font(r, 8, False, GRAY)

    def add_para(self, text="", bold=False, color=BLACK, size=None, after=5):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        _run_font(r, size or 10.2, bold, color)
        return p

    def add_bullet(self, text: str) -> None:
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        _run_font(r, 10.2)

    def add_number(self, text: str) -> None:
        p = self.doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        _run_font(r, 10.2)

    def add_callout(self, label: str, text: str, color=BLUE, fill=LIGHT_BLUE) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        _shading(c, fill)
        _margins(c, 130, 150, 130, 150)
        p = c.paragraphs[0]
        r = p.add_run(label + "  ")
        _run_font(r, 10, True, color)
        r2 = p.add_run(text)
        _run_font(r2, 10, False, BLACK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def add_table(self, headers, rows, widths=None, font=8.4):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        if widths:
            for i, w in enumerate(widths):
                t.columns[i].width = Inches(w)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = str(h)
            _shading(c, NAVY)
            _margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_font(c, font, True, WHITE, WD_ALIGN_PARAGRAPH.CENTER)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
                _margins(cells[i])
                if ri % 2:
                    _shading(cells[i], PALE)
                _cell_font(cells[i], font)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return t

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _v(self, *keys, default="—"):
        for k in keys:
            v = self.p.get(k)
            if v not in (None, ""):
                return v
        return default

    def _risk(self) -> tuple[str, str]:
        nature = str(self._v("enterprise_nature", default=""))
        if "国有" in nature:
            return "低风险", "公开身份维度可按低风险＋简易程序管理（须完成征信/中登等）。"
        return "一般关注", "企业性质待结合非公开材料与征信进一步判断。"

    def build(self, out_path: Path) -> Path:
        risk, risk_note = self._risk()
        rating = self._v("latest_subject_rating", default="")
        rating_txt = (
            f"{rating}（{self._v('latest_subject_rating_organization', default='')} "
            f"{self._v('latest_subject_rating_date', default='')}）".strip()
            if rating and rating != "—"
            else "公开主体评级暂缺，建议跟踪评级/债券披露或征信续核"
        )
        ctrl = self._v("controller")
        nature = self._v("enterprise_nature")
        judicial = self.extras.get("judicial_rows") or []
        penalty_note = self.extras.get("penalty_note") or "财汇主体监管处罚检索见正文；归档日须保留信用中国等官方留痕。"
        shareholder_rows = self.extras.get("shareholder_rows") or []
        control_path = self.extras.get("control_path") or f"实际控制人/控制关系：{ctrl}"
        finance_note = self.extras.get("finance_note") or "财务三源细核（iFinD/Wind）可在 Agent 侧按 skill 续核后补表。"
        date_cn = datetime.now().strftime("%Y年%m月%d日")

        # Cover
        self.add_para("客户尽职调查 · KYC", bold=True, color=GOLD, size=10.5, after=16)
        title = self.doc.add_paragraph(style="Title")
        title.add_run(self.name)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run("KYC 报告（2026三源核验版）")
        _run_font(r, 22, True, NAVY)
        self.add_para("依据《KYC工作指引（2026年修订版）》编制 · auto-kyc Skill", color=GRAY, size=11.5, after=28)
        self.add_table(
            ["项目", "内容"],
            [
                ["客户名称", self.name],
                ["法定代表人", self._v("legal_representative")],
                ["注册资本", str(self._v("registered_capital"))],
                ["成立日期", self._v("establishment_date")],
                ["注册地址", self._v("registered_address")],
                ["企业性质", nature],
                ["实际控制人", ctrl],
                ["主体评级（公开）", rating_txt],
                ["查询基准日", f"{self.checked_at}（{date_cn}）"],
                ["调查类型", "投前 KYC / 公开信息核验（Skill自动）"],
                ["数据来源", "财汇企业预警通 MCP；同花顺 iFinD / 万得 Wind 可续核"],
                ["保密级别", "内部资料"],
            ],
            [1.65, 4.8],
            9.5,
        )
        self.add_callout(
            "重要提示",
            "本报告为公开信息调查底稿（结构对齐上海城投三源核验版样例），不替代人行征信、证照原件、章程审阅、现场访谈及具体交易信用审批。",
            AMBER,
            "FFF2CC",
        )
        self.page_break()

        # 一、结论
        self.doc.add_heading("一、KYC综合结论", level=1)
        self.add_callout(
            "总体结论",
            f"客户身份风险建议认定为{risk}。{risk_note}最终授信/租赁决策以具体交易审批及非公开核验为准。",
            GREEN if risk == "低风险" else AMBER,
            "E2F0D9" if risk == "低风险" else "FFF2CC",
        )
        self.doc.add_heading("1. 风险等级与程序适用", level=2)
        self.add_para(
            f"{self.name}企业性质为「{nature}」，实际控制人/控制关系为「{ctrl}」。"
            f"依据《KYC工作指引（2026年修订版）》第十三条及附件1，"
            f"{'可适用简易程序前提：未见公开高风险直接认定事项。' if risk == '低风险' else '须结合非公开材料审慎适用程序分类。'}"
            "简易程序不免除官方信用信息、司法、监管及负面舆情核查。"
        )
        self.add_table(
            ["评价维度", "判断", "主要依据"],
            [
                ["客户特性", risk, f"企业性质 {nature}；实控人 {ctrl}；评级 {rating or '暂缺'}"],
                ["地域", "见注册地址", self._v("registered_address")],
                ["行业", "待项目化认定", "公开工商信息为主；主营细项建议 iFinD/Wind 续核"],
                ["业务/交易结构", "待项目化评估", "以立项申报书金额、期限、用途、增信为准"],
            ],
            [1.05, 1.05, 4.35],
            8.8,
        )
        self.doc.add_heading("2. 负面事项汇总与影响研判", level=2)
        neg_rows = self.extras.get("negative_rows") or [
            ["监管处罚", penalty_note[:80], "主体公开检索", "关注/复核"],
            [
                "司法案件",
                f"财汇公开检索返回 {len(judicial)} 条摘要（详见第六章）",
                "须核金额与执行状态",
                "关注进展" if judicial else "未见公开命中/待复核",
            ],
            [
                "公司负面舆情",
                (self.extras.get("news_note") or "舆情待财汇/公告续核")[:80],
                "与还款来源关联性待项目化判断",
                "须跟踪",
            ],
            ["主体评级", rating_txt, "公开评级字段", "一般关注" if not rating else "已披露"],
        ]
        self.add_table(["事项", "事实摘要", "关联性与风险影响", "结论"], neg_rows, [0.9, 2.15, 2.85, 0.65], 7.7)
        self.doc.add_heading("3. 准入意见与管控措施", level=2)
        for x in [
            f"建议附条件{'同意' if risk == '低风险' else '审慎'}准入：公开KYC维度为{risk}，最终以交易审批为准。",
            "项目申报前补齐营业执照、章程、法定代表人身份证明、银行账户信息，并与公示信息核对。",
            "取得人行征信及中登网完整登记，核验授信、担保、应收账款及融资租赁登记。",
            "明确本次业务用途、期限、还款来源和增信措施；核验债务到期分布。",
            "将评级变动、重大诉讼、对外担保、债券兑付纳入持续监测；iFinD/Wind 财务与债券口径按需续核。",
        ]:
            self.add_bullet(x)
        self.page_break()

        # 二、方法
        self.doc.add_heading("二、调查依据、范围与方法", level=1)
        self.doc.add_heading("1. 制度依据", level=2)
        self.add_para(
            "本报告依据《KYC工作指引（2026年修订版）》编制，落实“多源采集、交叉验证、分类研判、实质分析”。"
            "口径差异时以监管、司法及官方披露优先，第三方数据辅助。"
        )
        self.doc.add_heading("2. 数据源与交叉验证", level=2)
        self.add_table(
            ["数据源", "主要用途", "本次结果"],
            [
                ["财汇企业预警通 MCP", "工商、控制权、评级、司法、监管处罚", "已完成主体公开查询（见本报告）"],
                ["同花顺 iFinD", "财务、股东、公告与新闻", "可按 auto-kyc skill 在 Agent 侧续核"],
                ["万得 Wind", "档案、行业、财务、实控人交叉", "可按 auto-kyc skill 在 Agent 侧续核"],
                ["客户经理既有KYC/申报材料", "交叉核对", "平台已关联则见附件"],
            ],
            [1.45, 2.25, 2.75],
            8.6,
        )
        self.add_para(
            "口径差异提示：三源股东名次、财务快照日期可能不同；核心控制权与工商身份以官方一致结论为准。",
            color=GRAY,
            size=9.2,
        )
        self.doc.add_heading("3. 调查边界", level=2)
        for x in [
            "未替代人行征信、中登网全量登记及客户内部债务到期表。",
            "第三方库存在更新时滞与字段缺失；“未检索到”不等于法律上绝对不存在。",
            "本次为 Skill 自动公开KYC；完整三源细核可在 Agent 会话中续跑 iFinD/Wind。",
        ]:
            self.add_bullet(x)
        self.page_break()

        # 三、身份
        self.doc.add_heading("三、客户身份与行业认定", level=1)
        self.add_table(
            ["项目", "核验结果"],
            [
                ["企业全称", self.name],
                ["法定代表人", self._v("legal_representative")],
                ["注册资本", str(self._v("registered_capital"))],
                ["成立日期", self._v("establishment_date")],
                ["注册地址", self._v("registered_address")],
                ["企业性质", nature],
                ["实际控制人", ctrl],
                ["主体评级（公开）", rating_txt],
            ],
            [2.0, 4.45],
            9.1,
        )
        self.doc.add_heading("1. 主营业务与行业判断", level=2)
        self.add_para(
            self.extras.get("industry_note")
            or "公开工商信息已采集；实质行业与主营构成建议以 iFinD/Wind 及申报书业务描述交叉认定。"
        )
        self.doc.add_heading("2. 证照及身份文件核验清单", level=2)
        self.add_table(
            ["材料", "公开信息核验", "项目归档状态"],
            [
                ["营业执照", "工商关键信息已财汇核验", "待客户提供并核原件/电子证照"],
                ["公司章程", "未取得", "待补充"],
                ["法定代表人身份证明", "姓名已核验", "证件待补充"],
                ["银行账户信息", "未取得", "待补充"],
                ["人行征信/中登", "未取得", "放款前必须补齐"],
            ],
            [1.55, 2.55, 2.35],
            8.4,
        )
        self.page_break()

        # 四、股权
        self.doc.add_heading("四、股权结构、控制权与受益所有人", level=1)
        self.add_callout("控制路径", control_path, BLUE, LIGHT_BLUE)
        self.doc.add_heading("1. 控股股东与实际控制人", level=2)
        self.add_para(f"公开信息显示实际控制人/控制关系为「{ctrl}」。若三源股东表不一致，以最近一期官方披露为准。")
        self.doc.add_heading("2. 主要股东（公开快照）", level=2)
        if shareholder_rows:
            self.add_table(["序号", "股东名称", "持股比例", "备注"], shareholder_rows, [0.5, 3.55, 1.0, 1.4], 8.0)
        else:
            self.add_para("本次未结构化返回完整股东表；建议以财汇股权工具或 iFinD/Wind 前十大股东快照补表。")
        self.doc.add_heading("3. 受益所有人识别", level=2)
        self.add_para(
            "国有/机构实控情形下，按《指引》对自然人穿透阈值区别适用；本次以公开控制权为主。"
            "若出现代持或异常控制安排，应启动强化识别。"
        )
        self.page_break()

        # 五、财务
        self.doc.add_heading("五、财务状况与偿债能力观察", level=1)
        self.add_para(finance_note)
        self.add_callout("信用判断", "公开KYC阶段以身份、司法、评级与控制权为主；财务杠杆与现金流须在尽调中结合简易数据表与三源财报续核。", AMBER, "FFF2CC")
        self.page_break()

        # 六、司法监管
        self.doc.add_heading("六、经营预警、监管与司法核查", level=1)
        self.doc.add_heading("1. 监管处罚与失信", level=2)
        self.add_para(penalty_note)
        self.doc.add_heading("2. 司法案件", level=2)
        if judicial:
            self.add_table(
                ["日期/进程", "案由/摘要", "客户身份", "案号/备注"],
                judicial[:12],
                [1.05, 2.35, 0.7, 2.35],
                7.9,
            )
        else:
            self.add_para("财汇公开司法检索未返回可结构化明细，或工具字段为空；归档时须保留官方执行信息查询留痕。")
        self.doc.add_heading("3. 负面舆情", level=2)
        self.add_para(self.extras.get("news_note") or "舆情细核可结合财汇新闻/iFinD 公告检索在 Agent 侧续补。")
        self.page_break()

        # 七、交易
        self.doc.add_heading("七、本次业务交易背景调查", level=1)
        self.add_callout(
            "当前状态",
            "具体租赁要素以立项申报为准；公开KYC不能替代交易背景、租赁物权属与还款来源核验。",
            AMBER,
            "FFF2CC",
        )
        for x in [
            "核验承租人/担保人关系与交易合同链条。",
            "资金用途与主营、项目批复匹配，排除空转。",
            "租赁物权属、发票、验收及中登优先顺位。",
            "还款来源覆盖与增信可执行性。",
        ]:
            self.add_number(x)
        self.page_break()

        # 八、九、附录
        self.doc.add_heading("八、持续监测与复审安排", level=1)
        self.add_para("依《指引》第十九至二十一条，投后KYC与回访周期一致；重大负面、控制权变化、处罚、失信、破产、评级下调时随时审查。")
        self.doc.add_heading("九、最终结论", level=1)
        self.add_callout("建议", f"{risk}（公开KYC身份维度）；附条件准入。", GREEN if risk == "低风险" else AMBER, "E2F0D9" if risk == "低风险" else "FFF2CC")
        self.add_para(
            f"综合公开工商、控制权、评级与司法/处罚检索结果，{self.name} 公开维度建议按「{risk}」管理。"
            "最终准入须以征信、中登、交易背景及增信核验为前提。"
        )
        self.page_break()
        self.doc.add_heading("附录A：数据来源与口径说明", level=1)
        self.add_table(
            ["编号", "来源/底稿", "用途", "查询/披露时点"],
            [
                ["A1", "《KYC工作指引（2026年修订版）》", "制度依据", "本地制度文件"],
                ["A2", "财汇企业预警通 MCP", "工商/控制权/司法/处罚/评级", self.checked_at],
                ["A3", "同花顺 iFinD", "财务/股东/公告（可续核）", "Agent 侧按需"],
                ["A4", "万得 Wind", "档案/财务交叉（可续核）", "Agent 侧按需"],
                ["A5", "上海城投三源核验版样例", "报告结构对齐", "~/Downloads/KYC/"],
            ],
            [0.55, 2.15, 2.45, 1.3],
            8.2,
        )
        self.doc.add_heading("三源一致性对照", level=2)
        self.add_table(
            ["核验事项", "财汇", "iFinD", "Wind", "结论"],
            [
                ["企业名称", "已返回", "可续核", "可续核", "以财汇本次查询为准"],
                ["实控人/控制权", ctrl, "可续核", "可续核", "公开一致或待交叉"],
                ["主体评级", rating or "暂缺", "可续核", "可续核", "双口径披露"],
                ["司法/处罚", "已检索", "—", "—", "以财汇+官方归档复核"],
            ],
            [1.35, 1.05, 1.05, 1.05, 2.0],
            7.8,
        )
        self.doc.add_heading("重要限制与免责声明", level=2)
        self.add_para(
            "本报告基于查询基准日前公开或合法取得的数据，仅供内部KYC参考。不得替代法律意见、审计、评估、征信或监管证明。",
            color=GRAY,
            size=9.2,
        )
        self.doc.add_heading("附录B：项目经理归档清单", level=1)
        self.add_table(
            ["序号", "材料/动作", "状态", "责任人/日期"],
            [[str(i), m, "□待补 □完成", ""] for i, m in enumerate(
                [
                    "营业执照及统一社会信用代码核验",
                    "章程及授权机制核验",
                    "法定代表人身份证明",
                    "银行账户信息",
                    "人行征信报告",
                    "中登网全量查询",
                    "交易合同与租赁物权属",
                    "债务到期表与担保明细",
                    "国家企业信用/信用中国/执行信息留痕",
                    "投后监测频率与风险触发条款",
                ],
                1,
            )],
            [0.55, 3.75, 1.15, 1.1],
            8.2,
        )

        self.doc.core_properties.title = f"{self.name}KYC报告（2026三源核验版）"
        self.doc.core_properties.subject = "依据KYC工作指引（2026年修订版）· auto-kyc Skill"
        self.doc.core_properties.author = "auto-kyc"
        self.doc.core_properties.keywords = "KYC, 三源核验, 财汇, iFinD, Wind"
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out_path))
        return out_path


def build_sanyuan_report(
    company: str,
    profile: dict | None,
    extras: dict | None,
    out_path: Path,
    checked_at: str | None = None,
) -> Path:
    checked = checked_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return SanyuanBuilder(company, profile, extras, checked).build(out_path)
